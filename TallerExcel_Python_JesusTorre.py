import openpyxl
from tabulate import tabulate

def cargar_datos_excel(ruta_excel="personas.xlsx"):
    # Carga el archivo Excel y almacena sus registros en una lista de listas.
    # También se agrega un número de índice (#) para identificar cada registro.
    try:
        # Abre el archivo Excel y selecciona automáticamente la hoja activa.
        excel_dataframe = openpyxl.load_workbook(ruta_excel)
        dataframe = excel_dataframe.active

        data = []

        # Recorre las filas del archivo y construye cada registro temporalmente.
        for row in range(1, dataframe.max_row):
            # Cada registro comienza con un número consecutivo.
            _row = [row]
            # Recorre las columnas del archivo para obtener los datos de cada persona.
            for col in dataframe.iter_cols(1, dataframe.max_column):
                _row.append(col[row].value)
            # Guarda el registro completo dentro de la lista principal.
            data.append(_row)
        return data
    except FileNotFoundError:
        # Se ejecuta cuando el archivo Excel no existe en la ubicación indicada.
        print(f"Error: No se encontró el archivo '{ruta_excel}'.")
        return []
    except Exception as e:
        # Captura cualquier otro error inesperado durante la lectura del archivo.
        print(f"Ocurrió un error al leer el archivo: {e}")
        return []


def mostrar_tabla(datos, encabezados, alineacion=None):
    # Función auxiliar encargada de mostrar los datos en forma de tabla.
    if not datos:
        print("No hay datos para mostrar.")
        return

    # Si no se especifica una alineación, todas las columnas se centran.
    if alineacion is None:
        alineacion = ("center",) * len(encabezados)

    print(tabulate(datos, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion ))


# ==========================================
# FUNCIONES DE LOS RETOS PRÁCTICOS
# ==========================================

def mostrar_todos(datos):
    # Muestra todos los registros cargados desde el archivo Excel.
    encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
    mostrar_tabla(datos, encabezados)


def mostrar_nombre_email(datos):
    # Reto 1: Muestra únicamente el nombre y el correo electrónico.
    filtrados = [[fila[2], fila[4]] for fila in datos]

    encabezados = ["Name", "Email"]
    mostrar_tabla(filtrados, encabezados)


def contar_registros(datos):
    # Reto 2: Muestra la cantidad total de registros.

    # len() permite conocer cuántos registros contiene la lista.
    total = len(datos)

    print("=" * 30)
    print(f"📊 Cantidad total de registros: {total}")
    print("=" * 30)


def filtrar_por_empresa(datos):
    # Reto 3: Filtra las personas cuya empresa contiene 'Group'.

    # Define el texto que se utilizará para realizar el filtro.
    filtro = "Group"

    # Recorre los registros y conserva únicamente lo que tienen Group
    filtrados = [
        fila
        for fila in datos
        if fila[3] and filtro.lower() in str(fila[3]).lower()
    ]

    encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]

    print(f"Filtrando empresas que contienen: 'Group'")
    mostrar_tabla(filtrados, encabezados)


def buscar_por_id(datos):
    # Reto 4: Busca una persona utilizando su ID.

    # Solicita al usuario el ID que desea consultar.
    id_buscar = input( "Ingrese el ID de la persona a buscar: ").strip()

    # Compara el ID ingresado con el ID almacenado en cada registro.
    encontrados = [
        fila
        for fila in datos
        if str(fila[1]).strip() == id_buscar
    ]

    encabezados = ["#", "Id", "Name", "Company", "Email", "MAC Address"]

    if encontrados:
        # Si existe una coincidencia, muestra el registro encontrado.
        mostrar_tabla(encontrados, encabezados)
    else:
        # Si no existe, informa al usuario que no se encontró el ID.
        print(
            f"No se encontró ninguna persona con ese ID"
        )


# ==========================================
# RETO 5: MENÚ INTERACTIVO
# ==========================================

def menu_principal():
    # Muestra el menú principal y permite ejecutar los diferentes retos.

    # Carga los datos del archivo Excel una sola vez al iniciar el programa.
    datos = cargar_datos_excel("personas.xlsx")

    # Si no se pudieron cargar datos, finaliza el programa.
    if not datos:
        return

    # Mantiene el menú funcionando hasta que el usuario seleccione Salir.
    while True:
        print("=" * 40)
        print("   MENÚ INTERACTIVO - CONSULTA DE PERSONAS   ")
        print("=" * 40)
        print("1. Ver todos los registros")
        print("2. Ver solo Nombre y Email")
        print("3. Contar total de registros")
        print("4. Filtrar empresas con 'Group'")
        print("5. Buscar persona por ID")
        print("0. Salir")
        print("=" * 40)

        # Solicita al usuario la opción que desea ejecutar.
        opcion = input("Seleccione una opción: ").strip()

        # Ejecuta una función diferente dependiendo de la opción seleccionada.
        match opcion:
            case "1":
                mostrar_todos(datos)

            case "2":
                mostrar_nombre_email(datos)

            case "3":
                contar_registros(datos)

            case "4":
                filtrar_por_empresa(datos)

            case "5":
                buscar_por_id(datos)

            case "0":
                print("Sistea salido.")
                break

            case _:
                # Controla las opciones que no existen en el menú.
                print("Opción no válida.")


# Punto de entrada del programa.
# Solo ejecuta el menú cuando este archivo se ejecuta directamente.
if __name__ == "__main__":
    menu_principal()